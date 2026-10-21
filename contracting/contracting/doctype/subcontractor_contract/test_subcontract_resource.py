# Copyright (c) 2026, kazem and Contributors
# See license.txt

"""2026-10-18 BRD: Subcontract resource tab + Labor contract type.

Fixture work item W (from template T, qty 10):
    Material    1 x 100     Labor       1 x 50
    Subcontract 2 x 30      Equipment   1 x 20      -> 230 per unit
so a Tender built from it carries labor_cost 500, subcontract_cost 600,
equipment_cost 200, boq_direct_cost 2300.
"""

from io import BytesIO

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils import flt

from contracting.contracting.doctype.subcontractor_contract.subcontractor_contract import (
	get_unallocated_boq_items,
)
from contracting.contracting.doctype.subcontractor_contract.test_subcontractor_contract import (
	add_tender_line,
	ensure_resource_item,
	ensure_stock_item,
	make_draft_contract,
	make_resource_tender_and_project,
	submit_contract,
)
from contracting.contracting.doctype.tender.test_tender import ensure_customer, make_tender
from contracting.contracting.doctype.tender.tender import add_boq_item_from_template
from contracting.contracting.utils import allocation

SUBCONTRACT_GROUP = "Subcontractors Items"


def ensure_item(item_code, item_group):
	if not frappe.db.exists("Item", item_code):
		frappe.get_doc({
			"doctype": "Item", "item_code": item_code, "item_group": item_group,
			"stock_uom": "Nos", "is_stock_item": 0,
		}).insert(ignore_permissions=True)
	return item_code


def make_four_resource_template(name, subcontract_ratio=2):
	if frappe.db.exists("Work Item Template V2", name):
		return frappe.get_doc("Work Item Template V2", name)
	return frappe.get_doc({
		"doctype": "Work Item Template V2",
		"template_name": name,
		"uom": "Nos",
		"description": name + " description",
		"material_items": [{"item": ensure_stock_item("_Test SC Stock Material"), "qty_per_unit": 1, "rate": 100}],
		"labor_items": [{"item": ensure_resource_item("_Test SC Labor"), "qty_per_unit": 1, "rate": 50}],
		"subcontract_items": [{
			"item": ensure_item("_Test SC Subcontract", SUBCONTRACT_GROUP), "qty_per_unit": subcontract_ratio, "rate": 30,
		}],
		"equipment_items": [{"item": ensure_item("_Test SC Equipment Item", "Equipment"), "qty_per_unit": 1, "rate": 20}],
	}).insert(ignore_permissions=True)


def make_four_resource_tender(suffix, template="_SR Template", qty=10):
	"""A Tender whose second BOQ row W came from the four-resource template
	through the real apply path (add_boq_item_from_template, FR-05), and a
	Project pointing at it."""
	ensure_customer()
	make_four_resource_template(template)
	tender = make_tender("_sr_" + suffix)
	add_boq_item_from_template(tender.name, template, qty)
	tender.reload()
	project = frappe.get_doc({
		"doctype": "Project", "project_name": "_SR Project " + suffix, "tender": tender.name,
	}).insert(ignore_permissions=True)
	return project, tender, tender.boq_items[-1]


def rows_of(doc, fieldname, boq_row):
	return [r for r in doc.get(fieldname) if r.work_item == boq_row.idx]


class TestSubcontractTemplateAndTender(FrappeTestCase):
	def test_template_totals_include_subcontract(self):
		"""TC-01."""
		template = make_four_resource_template("_SR Template TC01")
		self.assertAlmostEqual(template.total_work_item_cost, 230.0, delta=1e-6)
		self.assertAlmostEqual(template.subcontract_items[0].amount, 60.0, delta=1e-6)

	def test_template_subcontract_item_group_enforced(self):
		"""TC-02."""
		template = frappe.get_doc({
			"doctype": "Work Item Template V2",
			"template_name": "_SR Template TC02",
			"uom": "Nos",
			"subcontract_items": [{"item": ensure_resource_item("_Test SC Labor"), "qty_per_unit": 1, "rate": 1}],
		})
		with self.assertRaisesRegex(frappe.ValidationError, SUBCONTRACT_GROUP):
			template.insert(ignore_permissions=True)

	def test_tender_propagation_and_subcontract_cost(self):
		"""TC-03."""
		_project, tender, boq_row = make_four_resource_tender("tc03")
		sub = rows_of(tender, "subcontract_items", boq_row)
		self.assertEqual(len(sub), 1)
		self.assertAlmostEqual(sub[0].qty, 20.0, delta=1e-6)
		self.assertAlmostEqual(sub[0].rate, 30.0, delta=1e-6)
		self.assertAlmostEqual(sub[0].direct_cost_amount, 600.0, delta=1e-6)
		self.assertAlmostEqual(boq_row.subcontract_cost, 600.0, delta=1e-6)
		self.assertAlmostEqual(boq_row.labor_cost, 500.0, delta=1e-6)
		self.assertAlmostEqual(boq_row.boq_direct_cost, 2300.0, delta=1e-6)

	def test_subcontract_rows_get_same_cost_layers_as_labor(self):
		"""TC-04: Safety Factor folds into a Subcontract row exactly as into
		a Labor row."""
		_project, tender, boq_row = make_four_resource_tender("tc04")
		tender.safety_factor_percent = 10
		tender.save(ignore_permissions=True)
		labor = rows_of(tender, "labor_items", boq_row)[0]
		sub = rows_of(tender, "subcontract_items", boq_row)[0]
		self.assertAlmostEqual(sub.safety_factor_amount, sub.direct_cost_amount * 0.1, delta=1e-6)
		self.assertAlmostEqual(
			sub.safety_factor_amount / sub.direct_cost_amount,
			labor.safety_factor_amount / labor.direct_cost_amount, delta=1e-9,
		)

	def test_cascade_delete_removes_subcontract_rows(self):
		"""TC-06."""
		_project, tender, boq_row = make_four_resource_tender("tc06")
		tender.set("boq_items", [r for r in tender.boq_items if r.name != boq_row.name])
		tender.save(ignore_permissions=True)
		for fieldname in ("material_items", "labor_items", "subcontract_items", "equipment_items"):
			self.assertFalse([r for r in tender.get(fieldname) if r.boq_row_id == boq_row.name], fieldname)

	def test_export_has_subcontract_sheet(self):
		"""TC-07."""
		from openpyxl import load_workbook

		from contracting.contracting.doctype.tender.tender import export_tender_xlsx

		_project, tender, _boq_row = make_four_resource_tender("tc07")
		export_tender_xlsx(tender.name)
		wb = load_workbook(BytesIO(frappe.response["filecontent"]))
		self.assertEqual(wb.sheetnames, ["BOQ Items", "Material", "Labor", "Subcontract", "Equipment"])
		self.assertEqual(wb["Subcontract"].max_row, 2)  # header + the TC-03 row


class TestLaborAndInstallingContracts(FrappeTestCase):
	def test_labor_contract_takes_only_labor(self):
		"""TC-08 / TC-09."""
		project, _tender, boq_row = make_four_resource_tender("tc08")
		contract = make_draft_contract(project, "Labor")
		add_tender_line(contract, boq_row, 4, 50)
		contract.insert(ignore_permissions=True)

		self.assertEqual(len(contract.contract_labor_items), 1)
		self.assertAlmostEqual(contract.contract_labor_items[0].qty, 4.0, delta=1e-6)
		for fieldname in ("contract_material_items", "contract_subcontract_items", "contract_equipment_items"):
			self.assertEqual(contract.get(fieldname), [], fieldname)
		self.assertAlmostEqual(flt(contract.contracted_items[0].tender_unit_cost), 50.0, delta=1e-6)
		self.assertEqual(contract.contracted_items[0].out_of_tender_scope, 0)

	def test_labor_contract_ceiling_and_red_flag(self):
		"""TC-09 + ceiling decision: 60 > labor 50/unit needs delegation,
		and is red once delegated."""
		project, _tender, boq_row = make_four_resource_tender("tc09")
		contract = make_draft_contract(project, "Labor")
		add_tender_line(contract, boq_row, 4, 60)
		with self.assertRaisesRegex(frappe.ValidationError, "ceiling"):
			contract.insert(ignore_permissions=True)

		contract.price_exception_acknowledged = 1
		contract.insert(ignore_permissions=True)
		self.assertEqual(contract.contracted_items[0].out_of_tender_scope, 1)

	def test_labor_contract_extra_row_item_rule(self):
		"""TC-10."""
		project, _tender, boq_row = make_four_resource_tender("tc10")
		contract = make_draft_contract(project, "Labor")
		add_tender_line(contract, boq_row, 4, 50)
		contract.append("contract_labor_items", {
			"work_item": boq_row.name, "item": ensure_resource_item("_Test SC Labor"), "qty": 1,
		})
		contract.insert(ignore_permissions=True)
		extra = contract.contract_labor_items[-1]
		self.assertEqual((extra.source, extra.out_of_tender_scope), ("Extra", 1))

		contract.append("contract_labor_items", {
			"work_item": boq_row.name, "item": ensure_item("_Test SC Subcontract", SUBCONTRACT_GROUP),
			"qty": 1, "rate": 1,
		})
		with self.assertRaisesRegex(frappe.ValidationError, "not allowed on this tab"):
			contract.save(ignore_permissions=True)

	def test_installing_contract_takes_labor_and_subcontract(self):
		"""TC-11, revised 2026-10-18: Installing claims Labor + Subcontract."""
		project, _tender, boq_row = make_four_resource_tender("tc11")
		contract = make_draft_contract(project, "Installing")
		add_tender_line(contract, boq_row, 4, 60)
		contract.insert(ignore_permissions=True)

		self.assertEqual(len(contract.contract_labor_items), 1)
		self.assertAlmostEqual(contract.contract_labor_items[0].qty, 4.0, delta=1e-6)
		self.assertEqual(len(contract.contract_subcontract_items), 1)
		self.assertAlmostEqual(contract.contract_subcontract_items[0].qty, 8.0, delta=1e-6)
		for fieldname in ("contract_material_items", "contract_equipment_items"):
			self.assertEqual(contract.get(fieldname), [], fieldname)
		# (labor 500 + subcontract 600) / 10
		self.assertAlmostEqual(flt(contract.contracted_items[0].tender_unit_cost), 110.0, delta=1e-6)

	def test_installing_ceiling_is_labor_plus_subcontract(self):
		project, _tender, boq_row = make_four_resource_tender("tc11c")
		contract = make_draft_contract(project, "Installing")
		add_tender_line(contract, boq_row, 4, 111)
		with self.assertRaisesRegex(frappe.ValidationError, "ceiling"):
			contract.insert(ignore_permissions=True)

	def test_subcontract_direct_commitment_and_cross_type(self):
		"""TC-12 / TC-13."""
		project, _tender, boq_row = make_four_resource_tender("tc12")
		first = make_draft_contract(project, "Installing")
		add_tender_line(first, boq_row, 4, 60)
		first.insert(ignore_permissions=True)
		submit_contract(first)

		# TC-12: labor 10 units, 4 committed; subcontract 20 units, 8 committed
		# -> 6 work-item units left on both.
		self.assertAlmostEqual(allocation.get_row_availability("Installing", boq_row.name), 6.0, delta=1e-6)
		second = make_draft_contract(project, "Installing")
		add_tender_line(second, boq_row, 7, 60)
		with self.assertRaisesRegex(frappe.ValidationError, "cannot allocate"):
			second.insert(ignore_permissions=True)

		# TC-13, revised 2026-10-20 (partial contracts): Supply and Install
		# can still take the untouched material / equipment lines in full, so
		# the work item is offered at the best line's 10 units.
		self.assertAlmostEqual(
			allocation.get_row_availability("Supplying and installing", boq_row.name), 10.0, delta=1e-6,
		)

	def test_supply_install_consumes_subcontract(self):
		"""TC-14 / TC-20."""
		project, _tender, boq_row = make_four_resource_tender("tc14")
		sai = make_draft_contract(project, "Supplying and installing")
		add_tender_line(sai, boq_row, 3, 200)
		sai.insert(ignore_permissions=True)
		# TC-20: all four tabs propagated.
		for fieldname, qty in (
			("contract_material_items", 3), ("contract_labor_items", 3),
			("contract_subcontract_items", 6), ("contract_equipment_items", 3),
		):
			self.assertEqual(len(sai.get(fieldname)), 1, fieldname)
			self.assertAlmostEqual(sai.get(fieldname)[0].qty, qty, delta=1e-6)
		submit_contract(sai)

		sub_row = rows_of(frappe.get_doc("Tender", boq_row.parent), "subcontract_items", boq_row)[0]
		self.assertAlmostEqual(
			allocation.get_resource_availability(allocation.SUBCONTRACT_DOCTYPE, sub_row.name), 14.0, delta=1e-6,
		)
		self.assertAlmostEqual(allocation.get_row_availability("Installing", boq_row.name), 7.0, delta=1e-6)

	def test_components_pricing_sums_four_tabs(self):
		"""TC-20: Components mode rolls up Material + Labor + Subcontract + Equipment."""
		project, _tender, boq_row = make_four_resource_tender("tc20")
		contract = make_draft_contract(project, "Supplying and installing")
		add_tender_line(contract, boq_row, 2, 0)
		contract.contracted_items[0].pricing_mode = "Components"
		for fieldname in ("contract_material_items", "contract_labor_items",
						  "contract_subcontract_items", "contract_equipment_items"):
			for row in contract.get(fieldname):
				row.rate = 1
		contract.insert(ignore_permissions=True)
		# (2 + 2 + 4 + 2) x 1 / 2
		self.assertAlmostEqual(contract.contracted_items[0].unit_price, 5.0, delta=1e-6)

	def test_components_line_needs_no_unit_price_total_line_does(self):
		"""unit_price is mandatory only outside Components mode: a Components
		line saves with it blank (the rollup fills it); a Total line doesn't."""
		project, _tender, boq_row = make_four_resource_tender("cmp_blank")
		contract = make_draft_contract(project, "Installing")
		add_tender_line(contract, boq_row, 4, None)
		contract.contracted_items[0].pricing_mode = "Components"
		for fieldname in ("contract_labor_items", "contract_subcontract_items"):
			for row in contract.get(fieldname):
				row.rate = 10
		contract.insert(ignore_permissions=True)
		# (labor 4 + subcontract 8) x 10 / 4
		self.assertAlmostEqual(contract.contracted_items[0].unit_price, 30.0, delta=1e-6)

		total = make_draft_contract(project, "Installing")
		add_tender_line(total, boq_row, 1, None)
		with self.assertRaisesRegex(frappe.ValidationError, "Unit Price is required"):
			total.insert(ignore_permissions=True)

	def test_work_item_without_the_types_resource_not_offered(self):
		"""FR-14 edge case (gap analysis C-5): a work item with no Labor rows
		has nothing for a Labor contract."""
		project, tender = make_resource_tender_and_project("sr_nolabor", qty=10)
		boq_row = tender.boq_items[0]
		tender.set("labor_items", [])
		tender.save(ignore_permissions=True)
		self.assertEqual(allocation.get_row_availability("Labor", boq_row.name), 0.0)
		offered = {r["work_item"] for r in get_unallocated_boq_items([tender.name], "Labor")}
		self.assertNotIn(boq_row.name, offered)
		self.assertAlmostEqual(allocation.get_row_availability("Equipment", boq_row.name), 10.0, delta=1e-6)

	def test_zero_ratio_subcontract_row_does_not_constrain(self):
		"""TC-24."""
		project, tender = make_resource_tender_and_project("sr_zero", qty=10)
		boq_row = tender.boq_items[0]
		tender.append("subcontract_items", {
			"work_item": boq_row.idx, "item": ensure_item("_Test SC Subcontract", SUBCONTRACT_GROUP),
			"qty_per_unit": 0, "rate": 30, "currency": "EGP",
		})
		tender.save(ignore_permissions=True)
		self.assertAlmostEqual(
			allocation.get_row_availability("Supplying and installing", boq_row.name), 10.0, delta=1e-6,
		)
		# Installing still has the fixture's Labor row (ratio 1) to claim.
		self.assertAlmostEqual(allocation.get_row_availability("Installing", boq_row.name), 10.0, delta=1e-6)


class TestInstallingSharesLabor(FrappeTestCase):
	def test_installing_and_labor_contracts_share_labor_rows(self):
		"""Labor rows committed by an Installing contract are gone for a Labor
		contract, and the reverse."""
		project, _tender, boq_row = make_four_resource_tender("share")
		installing = make_draft_contract(project, "Installing")
		add_tender_line(installing, boq_row, 4, 60)
		installing.insert(ignore_permissions=True)
		submit_contract(installing)
		self.assertAlmostEqual(allocation.get_row_availability("Labor", boq_row.name), 6.0, delta=1e-6)

		labor = make_draft_contract(project, "Labor")
		add_tender_line(labor, boq_row, 5, 50)
		labor.insert(ignore_permissions=True)
		submit_contract(labor)
		# labor 10 - 4 - 5 = 1 left; subcontract still 12 units (6 work-item
		# units) - the best line decides what's offered (2026-10-20).
		self.assertAlmostEqual(allocation.get_row_availability("Installing", boq_row.name), 6.0, delta=1e-6)
		labor_row = allocation.get_resource_rows(allocation.LABOR_DOCTYPE, allocation.get_work_item(boq_row.name))[0]
		self.assertAlmostEqual(
			allocation.get_resource_availability(allocation.LABOR_DOCTYPE, labor_row.name), 1.0, delta=1e-6,
		)

	def test_retype_keeps_labor_drops_subcontract(self):
		"""Installing -> Labor keeps the Labor rows and clears Subcontract."""
		project, _tender, boq_row = make_four_resource_tender("retype")
		contract = make_draft_contract(project, "Installing")
		add_tender_line(contract, boq_row, 4, 50)
		contract.insert(ignore_permissions=True)
		contract.type_subcontractor = "Labor"
		# 2026-10-20: the dropped Subcontract share has to go somewhere - the
		# remaining components must still add up to the line total.
		contract.contract_labor_items[0].rate = 50
		contract.save(ignore_permissions=True)
		self.assertEqual(len(contract.contract_labor_items), 1)
		self.assertEqual(contract.contract_subcontract_items, [])


class TestSubcontractPaymentConditionsAndAddendum(FrappeTestCase):
	def _installing_contract(self, suffix):
		project, _tender, boq_row = make_four_resource_tender(suffix)
		contract = make_draft_contract(project, "Installing")
		add_tender_line(contract, boq_row, 4, 60)
		contract.insert(ignore_permissions=True)
		return contract, boq_row

	def test_category_gating_on_installing(self):
		"""TC-18."""
		contract, boq_row = self._installing_contract("tc18")
		sub_row = contract.contract_subcontract_items[0]
		contract.append("payment_conditions", {
			"work_item": boq_row.name, "condition": "Sub", "percent": 50,
			"category": "Subcontract", "subcontract_item": sub_row.name,
		})
		contract.append("payment_conditions", {
			"work_item": boq_row.name, "condition": "Lab", "percent": 50, "category": "Labor",
		})
		contract.save(ignore_permissions=True)

		contract.payment_conditions[-1].category = "Material"
		with self.assertRaisesRegex(frappe.ValidationError, "not applicable to a Installing contract"):
			contract.save(ignore_permissions=True)

	def test_category_link_mismatch(self):
		"""TC-19."""
		contract, boq_row = self._installing_contract("tc19")
		contract.append("payment_conditions", {
			"work_item": boq_row.name, "condition": "Sub", "percent": 100,
			"category": "Labor", "subcontract_item": contract.contract_subcontract_items[0].name,
		})
		with self.assertRaisesRegex(frappe.ValidationError, "Category must be Subcontract"):
			contract.save(ignore_permissions=True)

	def test_addendum_on_installing_uses_subcontract_math(self):
		"""TC-21."""
		contract, boq_row = self._installing_contract("tc21")
		submit_contract(contract)

		def addendum(qty):
			doc = frappe.get_doc({
				"doctype": "Contractor Contract Addendum",
				"naming_series": "Subcontractor-ADDENDUM-.YYYY.-",
				"subcontractor_contract": contract.name,
			})
			doc.append("addendum_items", {"action": "Add", "work_item": boq_row.name, "qty_delta": qty, "unit_price": 60})
			doc.insert(ignore_permissions=True)
			return doc

		# 20 subcontract units, 8 committed: +2 work-item units (4 units) fits.
		addendum(2).submit()
		contract.reload()
		self.assertEqual(len(contract.contracted_items), 2)
		self.assertAlmostEqual(sum(flt(r.qty) for r in contract.contract_subcontract_items), 12.0, delta=1e-6)
		self.assertAlmostEqual(sum(flt(r.qty) for r in contract.contract_labor_items), 6.0, delta=1e-6)

		# 4 labor / 8 subcontract units left = 4 work-item units; 5 is over.
		with self.assertRaises(frappe.ValidationError):
			addendum(5).submit()
