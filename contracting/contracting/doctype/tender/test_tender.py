# Copyright (c) 2026, kazem and Contributors
# See license.txt

import frappe
from frappe.tests.utils import FrappeTestCase

from contracting.contracting.api import WORK_ITEM_GROUP, create_work_item_group

TEST_CUSTOMER = "_Test Tender Customer"


def ensure_customer():
	if not frappe.db.exists("Customer", TEST_CUSTOMER):
		frappe.get_doc({
			"doctype": "Customer",
			"customer_name": TEST_CUSTOMER,
			"customer_type": "Company",
		}).insert(ignore_permissions=True)
	return TEST_CUSTOMER


def make_template(name, uom="Nos", description="template reference text"):
	if frappe.db.exists("Work Item Template V2", name):
		frappe.delete_doc("Work Item Template V2", name, force=True, ignore_permissions=True)
	if frappe.db.exists("Item", "WI-" + name):
		frappe.delete_doc("Item", "WI-" + name, force=True, ignore_permissions=True)
	return frappe.get_doc({
		"doctype": "Work Item Template V2",
		"template_name": name,
		"uom": uom,
		"description": description,
	}).insert(ignore_permissions=True)


def make_project_tender(project_name, client=TEST_CUSTOMER):
	"""Tender.project_tender is mandatory - every Tender needs one of these
	to link to, even tests that don't care about Project Tender's own
	behaviour (see the Won Project Tender BRD, 2026-08-28)."""
	return frappe.get_doc({
		"doctype": "Project Tender",
		"naming_series": "PT-.YYYY.-",
		"project_name": project_name,
		"client": client,
	}).insert(ignore_permissions=True)


def make_tender(project_name, template=None, qty=10, description="per-tender scope", client=TEST_CUSTOMER):
	project_tender = make_project_tender(project_name, client=client)
	tender = frappe.get_doc({
		"doctype": "Tender",
		"naming_series": "TND-.YYYY.-",
		"project_tender": project_tender.name,
		"status": "Draft",
		"blanket_margin_percent": 20,
	})
	row = {"item_name": "Work item", "original_quantity": qty, "description": description}
	if template:
		row["work_item_template"] = template
	tender.append("boq_items", row)
	tender.insert(ignore_permissions=True)
	return tender


class TestTender(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		create_work_item_group()
		ensure_customer()
		if not frappe.db.exists("Item", "_Test Tender Labor"):
			frappe.get_doc({
				"doctype": "Item",
				"item_code": "_Test Tender Labor",
				"item_group": "Labor",
				"stock_uom": "Nos",
				"is_stock_item": 0,
			}).insert(ignore_permissions=True)

	# ---------- Work Item Template V2 -> Item ----------

	def test_template_creates_item(self):
		template = make_template("_Test WIT Create", uom="Cubic Meter", description="slab desc")

		self.assertEqual(template.item, "WI-_Test WIT Create")
		item = frappe.get_doc("Item", template.item)
		self.assertEqual(item.item_group, WORK_ITEM_GROUP)
		self.assertEqual(item.stock_uom, "Cubic Meter")
		self.assertEqual(item.is_stock_item, 0)
		self.assertEqual(item.is_sales_item, 1)
		self.assertEqual(item.description, "slab desc")

	def test_template_edit_syncs_item(self):
		template = make_template("_Test WIT Sync")

		template.description = "revised"
		template.uom = "Cubic Meter"
		template.save(ignore_permissions=True)

		item = frappe.get_doc("Item", template.item)
		self.assertEqual(item.description, "revised")
		self.assertEqual(item.stock_uom, "Cubic Meter")

	def test_template_rename_renames_item(self):
		template = make_template("_Test WIT Rename")
		if frappe.db.exists("Item", "WI-_Test WIT Renamed"):
			frappe.delete_doc("Item", "WI-_Test WIT Renamed", force=True, ignore_permissions=True)

		frappe.rename_doc("Work Item Template V2", template.name, "_Test WIT Renamed")

		self.assertFalse(frappe.db.exists("Item", "WI-_Test WIT Rename"))
		self.assertTrue(frappe.db.exists("Item", "WI-_Test WIT Renamed"))
		self.assertEqual(
			frappe.db.get_value("Work Item Template V2", "_Test WIT Renamed", "item"),
			"WI-_Test WIT Renamed",
		)

	def test_template_blocked_when_item_code_taken(self):
		"""An Item under our code that isn't ours must never be adopted onto
		a template - it would end up on a customer's Sales Order."""
		if not frappe.db.exists("Item", "WI-_Test WIT Squatted"):
			frappe.get_doc({
				"doctype": "Item",
				"item_code": "WI-_Test WIT Squatted",
				"item_group": "Civil",
				"stock_uom": "Nos",
				"is_stock_item": 0,
			}).insert(ignore_permissions=True)

		with self.assertRaises(frappe.ValidationError):
			frappe.get_doc({
				"doctype": "Work Item Template V2",
				"template_name": "_Test WIT Squatted",
				"uom": "Nos",
			}).insert(ignore_permissions=True)

	# ---------- Tender BOQ row ----------

	def test_boq_fetches_item_and_description(self):
		template = make_template("_Test WIT Fetch", description="TEMPLATE TEXT")
		tender = make_tender("_Test Fetch", template=template.name, description=None)

		row = tender.boq_items[0]
		self.assertEqual(row.item_code, template.item)
		self.assertEqual(row.description, "TEMPLATE TEXT")

	def test_boq_description_is_not_clobbered(self):
		"""fetch_if_empty: the template seeds a starting point, but the
		per-tender wording must survive every later save."""
		template = make_template("_Test WIT Keep", description="TEMPLATE TEXT")
		tender = make_tender("_Test Keep", template=template.name)

		tender.boq_items[0].description = "PER-TENDER TEXT"
		tender.save(ignore_permissions=True)
		tender.reload()

		self.assertEqual(tender.boq_items[0].description, "PER-TENDER TEXT")

	# ---------- Resource row amount (Direct + Safety Factor + Indirect Cost) ----------

	def test_amount_includes_indirect_cost(self):
		"""2026-08-28 BRD: a resource row's amount is Direct + Safety +
		Indirect, all three layers - not just the first two."""
		tender = make_tender("_Test Indirect Included", qty=1)
		tender.safety_factor_percent = 10
		tender.propagated_addition_percent = 20
		tender.append("labor_items", {
			"work_item": tender.boq_items[0].idx,
			"item": "_Test Tender Labor",
			"qty_per_unit": 1,
			"rate": 100,
		})
		tender.save(ignore_permissions=True)

		row = tender.labor_items[0]
		self.assertAlmostEqual(row.direct_cost_amount, 100)
		self.assertAlmostEqual(row.safety_factor_amount, 10)
		self.assertAlmostEqual(row.indirect_cost_amount, 22)
		self.assertAlmostEqual(row.amount, 132)

	def test_zero_indirect_cost_amount_unchanged(self):
		"""No propagated_addition_percent (no linked Project Tender
		addition) -> amount is Direct + Safety only, same as before this
		BRD - indirect cost is legitimately zero, not silently dropped."""
		tender = make_tender("_Test Indirect Zero", qty=1)
		tender.safety_factor_percent = 10
		tender.append("labor_items", {
			"work_item": tender.boq_items[0].idx,
			"item": "_Test Tender Labor",
			"qty_per_unit": 1,
			"rate": 100,
		})
		tender.save(ignore_permissions=True)

		row = tender.labor_items[0]
		self.assertAlmostEqual(row.indirect_cost_amount, 0)
		self.assertAlmostEqual(row.amount, 110)


	# ---------- Per-unit vs lump-sum fixed additions (2026-10-13 BRD) ----------

	def make_per_unit_tender(self, project_name, qty):
		"""TC-07: pull template T (resources 100/unit, 10%, Fixed 5 per
		unit, Fixed 50 lump) through the server-side append path."""
		from contracting.contracting.doctype.tender.tender import add_boq_item_from_template
		from contracting.contracting.doctype.work_item_template_v2.test_work_item_template_v2 import (
			brd_additions,
			make_priced_template,
		)

		template = make_priced_template("_Test WIT Per Unit Tender", brd_additions())
		tender = make_tender(project_name, qty=1)
		tender.boq_items = []
		tender.save(ignore_permissions=True)
		add_boq_item_from_template(tender.name, template.name, original_quantity=qty)
		tender.reload()
		return tender

	def test_pull_sets_per_unit_and_lump_sum(self):
		"""TC-07: additions = 1000 x 10% + 5 x 10 + 50 = 200."""
		tender = self.make_per_unit_tender("_Test Per Unit Pull", qty=10)
		row = tender.boq_items[0]

		self.assertAlmostEqual(row.other_additions_pct, 10)
		self.assertAlmostEqual(row.fixed_additions_per_unit, 5)
		self.assertAlmostEqual(row.fixed_additions, 50)
		self.assertAlmostEqual(row.boq_direct_cost, 1000 + 200)
		self.assertAlmostEqual(tender.total_additions, 200)

	def test_quantity_change_rescales_per_unit_only(self):
		"""TC-05: qty 10 -> 20 gives 2000 x 10% + 5 x 20 + 50 = 350."""
		tender = self.make_per_unit_tender("_Test Per Unit Rescale", qty=10)
		tender.boq_items[0].original_quantity = 20
		tender.save(ignore_permissions=True)
		tender.reload()

		self.assertAlmostEqual(tender.boq_items[0].boq_direct_cost, 2000 + 350)
		self.assertAlmostEqual(tender.total_additions, 350)

	def test_zero_quantity_keeps_lump_sum(self):
		"""TC-08: no quantity -> no per-unit contribution, lump sum stays."""
		tender = make_tender("_Test Per Unit Zero Qty", qty=0)
		tender.boq_items[0].fixed_additions_per_unit = 5
		tender.boq_items[0].fixed_additions = 50
		tender.save(ignore_permissions=True)

		self.assertAlmostEqual(tender.boq_items[0].boq_direct_cost, 50)
		self.assertAlmostEqual(tender.total_additions, 50)

	def test_legacy_lump_sum_row_unchanged(self):
		"""TC-09: a row with only fixed_additions (as every pre-deploy row
		has) costs exactly what it did - fixed_additions_per_unit defaults
		to 0."""
		tender = make_tender("_Test Per Unit Legacy", qty=10)
		tender.boq_items[0].fixed_additions = 50
		tender.append("labor_items", {
			"work_item": tender.boq_items[0].idx,
			"item": "_Test Tender Labor",
			"qty_per_unit": 1,
			"rate": 100,
		})
		tender.save(ignore_permissions=True)
		tender.reload()

		self.assertEqual(tender.boq_items[0].fixed_additions_per_unit, 0)
		self.assertAlmostEqual(tender.boq_items[0].boq_direct_cost, 1000 + 50)
		self.assertAlmostEqual(tender.total_additions, 50)


	# ---------- Client rate pins the sell rate (BOQ intake, 2026-10-21) ----------

	def _priced_tender(self, name, labor_rate=100):
		tender = make_tender(name, qty=10)
		tender.append("labor_items", {"work_item": 1, "item": "_Test Tender Labor", "qty_per_unit": 1,
									  "rate": labor_rate, "currency": tender.base_currency()})
		tender.save(ignore_permissions=True)
		return tender

	def test_client_rate_pins_sell_rate_and_derives_margin(self):
		tender = self._priced_tender("_Test client rate")
		row = tender.boq_items[0]
		cost = row.effective_unit_price
		self.assertTrue(cost > 0)
		row.client_rate = cost * 1.25
		tender.save(ignore_permissions=True)
		row = tender.boq_items[0]
		self.assertAlmostEqual(row.sell_rate, cost * 1.25)
		self.assertAlmostEqual(row.margin_percent, 25)
		self.assertAlmostEqual(row.sell_amount, cost * 1.25 * 10)

	def test_client_rate_below_cost_gives_negative_margin(self):
		tender = self._priced_tender("_Test client rate low")
		row = tender.boq_items[0]
		row.client_rate = row.effective_unit_price / 2
		tender.save(ignore_permissions=True)
		self.assertAlmostEqual(tender.boq_items[0].margin_percent, -50)

	def test_client_rate_with_no_cost_sells_at_client_rate(self):
		tender = make_tender("_Test client rate zero")
		tender.boq_items[0].client_rate = 300
		tender.save(ignore_permissions=True)
		row = tender.boq_items[0]
		self.assertEqual((row.sell_rate, row.margin_percent, row.sell_amount), (300, 0, 3000))

	def test_clearing_client_rate_returns_to_manual_margin(self):
		tender = self._priced_tender("_Test client rate clear")
		row = tender.boq_items[0]
		row.client_rate = row.effective_unit_price * 2
		tender.save(ignore_permissions=True)
		row = tender.boq_items[0]
		row.client_rate = 0
		row.margin_percent = 10
		tender.save(ignore_permissions=True)
		row = tender.boq_items[0]
		self.assertAlmostEqual(row.sell_rate, row.effective_unit_price * 1.1)


class TestTenderExport(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		create_work_item_group()
		ensure_customer()
		if not frappe.db.exists("Item", "_Test Tender Labor"):
			frappe.get_doc({
				"doctype": "Item",
				"item_code": "_Test Tender Labor",
				"item_group": "Labor",
				"stock_uom": "Nos",
				"is_stock_item": 0,
			}).insert(ignore_permissions=True)

	def tearDown(self):
		frappe.set_user("Administrator")

	def _tender(self, name="_Test Export"):
		tender = make_tender(name, qty=10, description="=SUM(A1)")
		tender.append("labor_items", {
			"work_item": tender.boq_items[0].idx,
			"item": "_Test Tender Labor",
			"qty_per_unit": 1,
			"rate": 100,
		})
		tender.append("labor_items", {"work_item": 1, "item": "_Test Tender Labor", "qty_per_unit": 1, "rate": 5})
		tender.save(ignore_permissions=True)
		# save() prunes orphans, so point the second row at a missing BOQ idx
		# behind its back (legacy / imported data).
		frappe.db.set_value("Tender Labor Item", tender.labor_items[1].name, "work_item", 99)
		tender.reload()
		return tender

	def _export(self, tender, scope="Whole Tender"):
		from io import BytesIO

		from openpyxl import load_workbook

		from contracting.contracting.doctype.tender.tender import export_tender_xlsx

		export_tender_xlsx(tender.name, scope)
		self.assertEqual(frappe.response["type"], "binary")
		return frappe.response["filename"], load_workbook(BytesIO(frappe.response["filecontent"]))

	def test_whole_export_sheets_and_columns(self):
		tender = self._tender()
		filename, wb = self._export(tender)

		self.assertEqual(filename, f"{tender.name}.xlsx")
		self.assertEqual(wb.sheetnames, ["BOQ Items", "Material", "Labor", "Subcontract", "Equipment"])
		self.assertEqual(wb["BOQ Items"].max_row, 2)
		self.assertEqual(wb["Labor"].max_row, 3)
		self.assertEqual(wb["Equipment"].max_row, 1)  # header only
		headers = [c.value for c in wb["Labor"][1]]
		self.assertEqual(headers[:2], ["#", "Work Item"])
		self.assertNotIn("boq_row_id", headers)
		self.assertTrue(wb["Labor"]["A1"].font.bold)

	def test_work_item_name_and_orphan(self):
		tender = self._tender()
		_f, wb = self._export(tender, "Labor")
		ws = wb["Labor"]
		self.assertEqual(ws["B2"].value, tender.boq_items[0].item_name)
		self.assertIsNone(ws["B3"].value)

	def test_single_tab(self):
		tender = self._tender()
		filename, wb = self._export(tender, "Labor")
		self.assertEqual(filename, f"{tender.name} - Labor.xlsx")
		self.assertEqual(wb.sheetnames, ["Labor"])

	def test_formula_text_stays_text(self):
		tender = self._tender()
		_f, wb = self._export(tender, "BOQ Items")
		ws = wb["BOQ Items"]
		cells = [c for c in ws[2] if c.value == "=SUM(A1)"]
		self.assertTrue(cells)
		self.assertEqual(cells[0].data_type, "s")

	def test_disallowed_role_and_bad_scope(self):
		from contracting.contracting.doctype.tender.tender import export_tender_xlsx

		tender = self._tender()
		with self.assertRaises(frappe.ValidationError):
			export_tender_xlsx(tender.name, "Foo")

		email = "_test_export_denied@example.com"
		if not frappe.db.exists("User", email):
			frappe.get_doc({
				"doctype": "User", "email": email, "first_name": "Denied",
				"send_welcome_email": 0, "roles": [{"role": "Contracts Manager"}],
			}).insert(ignore_permissions=True)
		frappe.set_user(email)
		with self.assertRaises(frappe.PermissionError):
			export_tender_xlsx(tender.name, "Whole Tender")

	def test_export_has_no_side_effects(self):
		tender = self._tender()
		modified = frappe.db.get_value("Tender", tender.name, "modified")
		self._export(tender)
		self.assertEqual(frappe.db.get_value("Tender", tender.name, "modified"), modified)

	def test_grant_patch_idempotent(self):
		from contracting.patches import grant_tender_export_roles

		def rows():
			return frappe.db.count("Custom DocPerm", {"parent": "Tender"})

		grant_tender_export_roles.execute()
		first = rows()
		grant_tender_export_roles.execute()
		self.assertEqual(rows(), first)
		perm = frappe.db.get_value(
			"Custom DocPerm",
			{"parent": "Tender", "role": "Business Development Manager"},
			["read", "write", "create"], as_dict=True,
		)
		self.assertEqual((perm.read, perm.write, perm.create), (1, 0, 0))
